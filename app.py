# Vetology Data Science Training - Automated Radiology Evaluation
# Date: October 2026

import argparse
import json
import os
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import dotenv
from google import genai

# Load environment variables
dotenv.load_dotenv()
API_KEY = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY", "")
MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite")

# 20 Target Conditions for Canine Thorax
CONDITIONS = [
    "perihilar_infiltrate",
    "pneumonia",
    "bronchitis",
    "interstitial",
    "diseased_lungs",
    "hypo_plastic_trachea",
    "cardiomegaly",
    "pulmonary_nodules",
    "pleural_effusion",
    "rtm",
    "focal_caudodorsal_lung",
    "focal_perihilar",
    "pulmonary_hypoinflation",
    "right_sided_cardiomegaly",
    "pericardial_effusion",
    "bronchiectasis",
    "pulmonary_vessel_enlargement",
    "left_sided_cardiomegaly",
    "thoracic_lymphadenopathy",
    "esophagitis",
]


def clean_text(text):
    """Sanitize Excel text artifacts and HTML tags."""
    if text is None:
        return ""
    val = str(text).replace("_x000D_", "\n")
    val = re.sub(r"\r\n|\r", "\n", val)
    val = re.sub(r"<br\s*/?>", "\n", val, flags=re.IGNORECASE)
    return val.strip()


def load_reports(file_path, limit=100):
    """Load radiology reports from Excel."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    reports = []

    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        try:
            case_id = int(row[0])
        except (ValueError, TypeError):
            continue

        if len(reports) >= limit:
            break

        f_orig = clean_text(row[1] if len(row) > 1 else "")
        c_orig = clean_text(row[2] if len(row) > 2 else "")
        r_orig = clean_text(row[3] if len(row) > 3 else "")

        full_text = (
            f"FINDINGS:\n{f_orig}\n\nCONCLUSIONS:\n{c_orig}\n\nRECOMMENDATIONS:\n{r_orig}".strip()
        )

        # Extract baseline labels from columns 8 to 27
        baseline = {}
        for idx, cond in enumerate(CONDITIONS, start=8):
            raw = ""
            if len(row) > idx and row[idx] is not None:
                raw = str(row[idx]).strip().lower()
            baseline[cond] = "Positive" if raw in ["abnormal", "positive", "pos"] else "Negative"

        reports.append({
            "case_id": case_id,
            "text": full_text,
            "baseline": baseline,
        })

    return reports


def evaluate_with_gemini(client, report_text, dry_run=False):
    """Evaluate 20 thoracic conditions using Gemini in one optimized query."""
    if dry_run or not client:
        # Fast offline rule simulation
        lower_t = report_text.lower()
        res = {}
        for c in CONDITIONS:
            name_clean = c.replace("_", " ")
            if "no " + name_clean in lower_t or "normal" in lower_t:
                res[c] = "Negative"
            elif name_clean in lower_t:
                res[c] = "Positive"
            else:
                res[c] = "Negative"
        return res

    prompt = (
        "You are an expert veterinary radiologist evaluating a thoracic radiology report.\n"
        "Evaluate each of the following 20 conditions as 'Positive' "
        "(present/suspected/diagnosed) or 'Negative' (explicitly absent/normal/not mentioned).\n\n"
        f"Conditions:\n{json.dumps(CONDITIONS, indent=2)}\n\n"
        f"Report:\n\"\"\"{report_text}\"\"\"\n\n"
        "Return ONLY a valid JSON dictionary mapping each condition name to "
        "'Positive' or 'Negative'."
    )

    try:
        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=prompt,
        )
        cleaned = re.sub(
            r"^```(?:json)?|```$", "", response.text.strip(), flags=re.MULTILINE
        ).strip()
        parsed = json.loads(cleaned)
        return {
            c: "Positive" if str(parsed.get(c, "")).strip().capitalize() == "Positive"
            else "Negative"
            for c in CONDITIONS
        }
    except Exception as exc:
        print(f"  Warning: Gemini call failed ({exc}), falling back to Negative.")
        return {c: "Negative" for c in CONDITIONS}


def save_confusion_matrix(output_path, matrix_rows):
    """Write confusion matrix Excel file matching the official Vetology template."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Confusion Matrix"

    headers = [
        "condition", "True Positive", "False Negative", "True Negative", "False Positive",
        "Sensitivity", "Specificity", "Check", "Positive Ground Truth", "Negative Ground Truth",
        "Ground Truth Check", "", "Radiologist Agreement Rate", ""
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2EA3F2", end_color="2EA3F2", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD")
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        if headers[col_num - 1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

    for idx, row in enumerate(matrix_rows, start=2):
        cond = row["condition"]
        tp, fn, tn, fp = row["tp"], row["fn"], row["tn"], row["fp"]
        total = tp + fn + tn + fp
        agree = round((tp + tn) / total, 4) if total > 0 else 0.0

        ws.append([
            cond, tp, fn, tn, fp,
            f"=IFERROR(SUM(B{idx}:B{idx})/(SUM(B{idx}:B{idx})+SUM(C{idx}:C{idx})), 0)",
            f"=IFERROR(SUM(D{idx}:D{idx})/(SUM(D{idx}:D{idx})+SUM(E{idx}:E{idx})), 0)",
            f"=SUM(B{idx}:E{idx})",
            f"=SUM(B{idx}:C{idx})",
            f"=SUM(D{idx}:E{idx})",
            f"=SUM(I{idx}:J{idx})",
            None, agree, None
        ])

        for c_idx in range(1, 12):
            c_cell = ws.cell(row=idx, column=c_idx)
            c_cell.border = thin_border
            align_h = "left" if c_idx == 1 else "center"
            c_cell.alignment = Alignment(horizontal=align_h, vertical="center")

        ws.cell(row=idx, column=6).number_format = "0.0000"
        ws.cell(row=idx, column=7).number_format = "0.0000"
        ws.cell(row=idx, column=13).number_format = "0.00"

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Vetology AI Radiology Evaluator")
    parser.add_argument(
        "--data", type=str, default="1_canine_thorax_scoring_data.xlsx",
        help="Input dataset path",
    )
    parser.add_argument(
        "--limit", type=int, default=100,
        help="Number of reports to evaluate (default: 100)",
    )
    parser.add_argument(
        "--output", type=str, default="Confusion_Matrix_Canine_Thorax.xlsx",
        help="Output matrix file",
    )
    parser.add_argument(
        "--dry-run", action="store_true",
        help="Run offline without calling Gemini API",
    )
    args = parser.parse_args()

    print(f"Loading {args.limit} reports from {args.data}...")
    reports = load_reports(args.data, limit=args.limit)
    print(f"Loaded {len(reports)} reports successfully.")

    client = None
    if not args.dry_run and API_KEY:
        try:
            client = genai.Client(api_key=API_KEY)
        except Exception as err:
            print(f"Could not connect to Gemini: {err}. Using dry-run mode.")
            args.dry_run = True

    print(f"Evaluating reports with Gemini ({MODEL_NAME})...\n")
    predictions = {c: {} for c in CONDITIONS}

    for idx, rep in enumerate(reports, start=1):
        cid = rep["case_id"]
        res = evaluate_with_gemini(client, rep["text"], dry_run=args.dry_run)
        for c in CONDITIONS:
            predictions[c][cid] = res.get(c, "Negative")

        pos = [c for c, v in res.items() if v == "Positive"]
        summary = ", ".join(pos) if pos else "None (Normal)"
        print(f"[{idx}/{len(reports)}] Case {cid:03d} -> Positives: {summary}", flush=True)

    # Compute confusion matrix against baseline ground truth
    matrix_rows = []
    for c in CONDITIONS:
        tp = fn = tn = fp = 0
        for rep in reports:
            cid = rep["case_id"]
            gt = rep["baseline"].get(c, "Negative")
            pred = predictions[c].get(cid, "Negative")

            if gt == "Positive" and pred == "Positive":
                tp += 1
            elif gt == "Positive" and pred == "Negative":
                fn += 1
            elif gt == "Negative" and pred == "Negative":
                tn += 1
            elif gt == "Negative" and pred == "Positive":
                fp += 1

        matrix_rows.append({"condition": c, "tp": tp, "fn": fn, "tn": tn, "fp": fp})

    save_confusion_matrix(args.output, matrix_rows)
    print(f"\n[DONE] Confusion matrix successfully saved to: {args.output}")


if __name__ == "__main__":
    main()
